"""Iteration 11: XLSX export endpoints (library, prices, orders) + CSV template regression."""
import io
import openpyxl
import pytest
from conftest import BASE_URL

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _open_wb(content: bytes):
    return openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)


def test_library_export_xlsx(api_client, auth_headers):
    r = api_client.get(f"{BASE_URL}/api/library/export", headers=auth_headers)
    assert r.status_code == 200, r.text[:500]
    assert XLSX_MIME in r.headers.get("Content-Type", ""), r.headers
    assert "cost_library.xlsx" in r.headers.get("Content-Disposition", "")
    wb = _open_wb(r.content)
    assert "Cost Library" in wb.sheetnames
    ws = wb["Cost Library"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    expected = ["SKU", "Product", "Production Cost (EUR)", "Operational Cost (EUR)",
                "Production Shipping Cost (EUR)", "Commission %", "Commission EUR / unit",
                "Total Cost (EUR)", "Units Sold", "Revenue (EUR)"]
    assert headers == expected, headers


def test_prices_export_xlsx(api_client, auth_headers):
    r = api_client.get(f"{BASE_URL}/api/skus/prices/export",
                       params={"date_from": "2026-01-01", "date_to": "2026-12-31"},
                       headers=auth_headers)
    assert r.status_code == 200, r.text[:500]
    assert XLSX_MIME in r.headers.get("Content-Type", "")
    assert ".xlsx" in r.headers.get("Content-Disposition", "")
    wb = _open_wb(r.content)
    assert "SKU Prices" in wb.sheetnames
    ws = wb["SKU Prices"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert headers[0] == "SKU"
    assert headers[1] == "Product"
    assert len(headers) > 2, f"expected marketplace columns: {headers}"


def test_orders_export_xlsx(api_client, auth_headers):
    r = api_client.get(f"{BASE_URL}/api/orders/export",
                       params={"date_from": "2026-06-01", "date_to": "2026-06-30"},
                       headers=auth_headers)
    assert r.status_code == 200, r.text[:500]
    assert XLSX_MIME in r.headers.get("Content-Type", "")
    assert ".xlsx" in r.headers.get("Content-Disposition", "")
    wb = _open_wb(r.content)
    assert "Orders" in wb.sheetnames
    ws = wb["Orders"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert len(headers) == 19, f"expected 19 cols, got {len(headers)}: {headers}"
    assert headers[:5] == ["order_date_iso", "marketplace", "channel_raw", "order_id", "sku"], headers[:5]


# Regression: input templates should remain CSV
def test_asin_mapping_template_still_csv(api_client, auth_headers):
    r = api_client.get(f"{BASE_URL}/api/templates/asin-mapping", headers=auth_headers)
    assert r.status_code == 200
    assert "text/csv" in r.headers.get("Content-Type", "")


def test_cost_template_still_csv(api_client, auth_headers):
    r = api_client.get(f"{BASE_URL}/api/templates/cost", headers=auth_headers)
    assert r.status_code == 200
    assert "text/csv" in r.headers.get("Content-Type", "")
